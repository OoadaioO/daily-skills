#!/usr/bin/env python3
"""sheet-json: JSON ⇄ Excel 双向转换与静态校验。

依赖：仅 openpyxl。不依赖 LibreOffice，不依赖公式计算引擎。

子命令：
  to-json   xlsx → sheet-json
  to-excel  sheet-json → xlsx（写入前做静态校验）
  check     sheet-json 或 xlsx → 静态校验报告（JSON）
  verify    xlsx → json → xlsx 往返比对

用法见 ../SKILL.md
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
import tempfile
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.formula import ArrayFormula
except ImportError:  # pragma: no cover
    sys.exit("缺少依赖：请用 .venv/bin/python 运行，或执行 .venv/bin/pip install openpyxl")

SCHEMA_ID = "sheet-json/v1"

# ---------------------------------------------------------------- 常量表

PRESETS: dict[str, str] = {
    "general": "General",
    "integer": "#,##0;(#,##0);-",
    "decimal": "#,##0.00;(#,##0.00);-",
    "currency": "$#,##0;($#,##0);-",
    "currency2": "$#,##0.00;($#,##0.00);-",
    "percent": "0.0%;(0.0%);-",
    "multiple": "0.0x;(0.0x);-",
    "text": "@",
    "date": "yyyy-mm-dd",
    "datetime": "yyyy-mm-dd hh:mm",
}
PRESET_BY_FORMAT = {v: k for k, v in PRESETS.items()}
T_DEFAULT_FORMAT = {"date": "yyyy-mm-dd", "datetime": "yyyy-mm-dd hh:mm", "time": "hh:mm:ss"}

# role → (字体 ARGB | None, 底色 ARGB | None)，对应 xlsx/SKILL.md 的颜色编码
ROLE_STYLE: dict[str, tuple[str | None, str | None]] = {
    "input": ("0000FF", None),
    "formula": ("000000", None),
    "link": ("008000", None),
    "external": ("FF0000", None),
    "assumption": (None, "FFFF00"),
}
ROLE_NAMES = set(ROLE_STYLE)
FONT_ROLE_BY_RGB = {"0000FF": "input", "008000": "link", "FF0000": "external"}
FILL_ROLE_BY_RGB = {"FFFF00": "assumption"}

EXCEL_ERRORS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A", "#GETTING_DATA")
# 公式串里出现即视为已损坏的错误字面量（#N/A 可能是 IF(ISNA(...), #N/A, ...) 的有意写法）
FATAL_ERROR_LITERALS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!")

MAX_COL, MAX_ROW = 16384, 1048576
CELL_OBJ_KEYS = {"v", "f", "t", "preset", "fmt", "role", "note", "link", "array"}
SHEET_KEYS = {"name", "header", "columns", "rows", "merges", "columnWidths", "freeze"}
BAD_SHEET_CHARS = set('[]:*?/\\')

# Excel 内建函数白名单。可经 --extra-funcs 追加。
EXCEL_FUNCTIONS = set("""
ABS ACCRINT ACCRINTM ACOS ACOSH ACOT ACOTH AGGREGATE ADDRESS AMORDEGRC AMORLINC AND
ARABIC AREAS ARRAYTOTEXT ASC ASIN ASINH ATAN ATAN2 ATANH AVEDEV AVERAGE AVERAGEA
AVERAGEIF AVERAGEIFS BAHTTEXT BASE BESSELI BESSELJ BESSELK BESSELY BETADIST BETAINV
BETA.DIST BETA.INV BIN2DEC BIN2HEX BIN2OCT BINOMDIST BINOM.DIST BINOM.DIST.RANGE
BINOM.INV BITAND BITLSHIFT BITOR BITRSHIFT BITXOR BYCOL BYROW CALL CEILING CEILING.MATH
CEILING.PRECISE CELL CHAR CHIDIST CHIINV CHITEST CHISQ.DIST CHISQ.DIST.RT CHISQ.INV
CHISQ.INV.RT CHISQ.TEST CHOOSE CHOOSECOLS CHOOSEROWS CLEAN CODE COLUMN COLUMNS COMBIN
COMBINA COMPLEX CONCAT CONCATENATE CONFIDENCE CONFIDENCE.NORM CONFIDENCE.T CONVERT CORREL
COS COSH COT COTH COUNT COUNTA COUNTBLANK COUNTIF COUNTIFS COUPDAYBS COUPDAYS
COUPDAYSNC COUPNCD COUPNUM COUPPCD COVAR COVARIANCE.P COVARIANCE.S CRITBINOM CSC CSCH
CUBEKPIMEMBER CUBEMEMBER CUBEMEMBERPROPERTY CUBERANKEDMEMBER CUBESET CUBESETCOUNT
CUBEVALUE CUMIPMT CUMPRINC DATE DATEDIF DATEVALUE DAVERAGE DBCS DCOUNT DCOUNTA DDB DEC2BIN
DEC2HEX DEC2OCT DECIMAL DEGREES DELTA DEVSQ DGET DISC DMAX DMIN DOLLAR DOLLARDE DOLLARFR
DPRODUCT DSTDEV DSTDEVP DSUM DURATION DVAR DVARP ECDF EDATE EFFECT ENCODEURL EOMONTH
ERF ERF.PRECISE ERFC ERFC.PRECISE ERROR.TYPE EUROCONVERT EVEN EXACT EXP EXPAND
EXPON.DIST EXPONDIST F.DIST F.DIST.RT F.INV F.INV.RT F.TEST FACT FACTDOUBLE FALSE F.DIST
FDIST F.DIST.RT FDIST F.INV FINDB FIND FINV FISHER FISHERINV FIXED FLOOR FLOOR.MATH
FLOOR.PRECISE FORECAST FORECAST.ETS FORECAST.ETS.CONFINT FORECAST.ETS.SEASONALITY
FORECAST.ETS.STAT FORECAST.LINEAR FORMULATEXT FREQUENCY F.TEST FTEST FV FVSCHEDULE
GAMMA GAMMA.DIST GAMMA.INV GAMMA.DIST GAMMADIST GAMMAINV GAMMALN GAMMALN.PRECISE GAUSS
GCD GEOMEAN GESTEP GETPIVOTDATA GROWTH HARMEAN HEX2BIN HEX2DEC HEX2OCT HLOOKUP HOUR
HSTACK HYPERLINK HYPGEOM.DIST HYPGEOMDIST IF IFERROR IFNA IFS IMABS IMAGINARY IMARGUMENT
IMCONJUGATE IMCOS IMCOSH IMCOT IMCSC IMCSCH IMDIV IMEXP IMLN IMLOG10 IMLOG2 IMPOWER
IMPRODUCT IMREAL IMSEC IMSECH IMSIN IMSINH IMSQRT IMSUB IMSUM IM TAN INDEX INDIRECT
INFO INT INTERCEPT INTRATE IPMT IRR ISBLANK ISERR ISERROR ISEVEN ISFORMULA ISLOGICAL
ISNA ISNONTEXT ISNUMBER ISODD ISREF ISTEXT ISO.CEILING ISOWEEKNUM ISPMT JIS KURT LARGE
LCM LEFT LEFTB LEN LENB LET LINEST LN LOG LOG10 LOGEST LOGINV LOGNORM.DIST LOGNORM.INV
LOGNORMDIST LOGNORMDIST LOOKUP LOWER MAKEARRAY MAP MATCH MAX MAXA MAXIFS MDETERM MDURATION
MEDIAN MID MIDB MIN MINA MINIFS MINUTE MINVERSE MIRR MMULT MOD MODE MODE.MULT MODE.SNGL
MONTH MROUND MULTINOMIAL MUNIT N NA NEGBINOM.DIST NEGBINOMDIST NETWORKDAYS
NETWORKDAYS.INTL NOMINAL NORM.DIST NORM.INV NORM.S.DIST NORM.S.INV NORMDIST NORMINV
NORMSDIST NORMSINV NOT NOW NPER NPV NUMBERVALUE OCT2BIN OCT2DEC OCT2HEX ODD ODDFPRICE
ODDFYIELD ODDLPRICE ODDLYIELD OFFSET OR PDURATION PEARSON PERCENTILE PERCENTILE.EXC
PERCENTILE.INC PERCENTRANK PERCENTRANK.EXC PERCENTRANK.INC PERMUT PERMUTATIONA PHI
PHONETIC PI PMT POISSON POISSON.DIST POWER PPMT PRICE PRICEDISC PRICEMAT PROB PRODUCT
PROPER PV QUARTILE QUARTILE.EXC QUARTILE.INC QUOTIENT RADIANS RAND RANDARRAY RANDBETWEEN
RANK RANK.AVG RANK.EQ RATE RECEIVED REDUCE REGEXEXTRACT REGEXREPLACE REGEXTEST
REGISTER.ID REPLACE REPLACEB REPT RIGHT RIGHTB ROMAN ROUND ROUNDDOWN ROUNDUP ROW ROWS
RRI RSQ RTD SEARCH SEARCHB SEC SECH SECOND SEQUENCE SERIESSUM SHEET SHEETS SIGN SIN SINH
SKEW SKEW.P SLN SLOPE SMALL SORT SORTBY SQRT SQRTPI STANDARDIZE STDEV STDEV.P STDEV.S
STDEVA STDEVP STDEVPA STEYX SUBSTITUTE SUBTOTAL SUM SUMIF SUMIFS SUMPRODUCT SUMSQ SUMX2MY2
SUMX2PY2 SUMXMY2 SWITCH SYD T TAN TANH TAKE TBILLEQ TBILLPRICE TBILLYIELD T.DIST
T.DIST.2T T.DIST.RT T.INV T.INV.2T T.TEST TAN TANH TEXT TEXTAFTER TEXTBEFORE TEXTJOIN
TEXTSPLIT TIME TIMEVALUE TINV TINV TOCOL TOROW TODAY TRANSPOSE TREND TRIM TRIMMEAN TRUE
TRUNC T.TEST TTEST TYPE UNICHAR UNICODE UNIQUE UPPER VALUE VALUETOTEXT VAR VAR.P VAR.S
VARA VARP VARPA VDB VLOOKUP VSTACK WEEKDAY WEEKNUM WEIBULL WEIBULL.DIST WORKDAY
WORKDAY.INTL WRAPCOLS WRAPROWS XIRR XLOOKUP XMATCH XOR YEAR YEARFRAC YIELD YIELDDISC
YIELDMAT Z.TEST ZTEST
""".split())


# ---------------------------------------------------------------- 通用工具


def die(msg: str, code: int = 2):
    print(f"错误：{msg}", file=sys.stderr)
    raise SystemExit(code)


def col_to_idx(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def parse_a1(ref: str) -> tuple[int, int]:
    """'$B$6' → (2, 6)"""
    m = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?(\d{1,7})", ref)
    if not m:
        raise ValueError(ref)
    return col_to_idx(m.group(1)), int(m.group(2))


def is_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


# ---------------------------------------------------------------- 公式分析

CELL_RE = re.compile(r"(?<![A-Za-z0-9_])(\$?[A-Za-z]{1,3}\$?\d{1,7})(?![A-Za-z0-9_(])")
FUNC_RE = re.compile(r"(?<![A-Za-z0-9_.])([A-Za-z_][A-Za-z0-9_.]*)\s*\(")
NUM_RE = re.compile(r"(?<![A-Za-z0-9_.])(\d+(?:\.\d+)?)(?![A-Za-z0-9_.])")
PREFIX_RE = re.compile(
    r"(?:\[([^\]]*)\])?"
    r"(?:'((?:[^']|'')+)'|([A-Za-z_\u4e00-\u9fff][^!'()+\-*/^&<>=,;:\[\]\s]*))!$"
)


def mask_formula(f: str) -> tuple[str, list[tuple[int, int]]]:
    """把字符串字面量替换为等长空格；返回 (掩码串, 引号内表名区间)。

    等长替换保证位置信息可用于过滤匹配。
    """
    chars = list(f)
    sheet_spans: list[tuple[int, int]] = []
    i, n = 0, len(f)
    while i < n:
        c = f[i]
        if c == '"':
            j = i + 1
            while j < n:
                if f[j] == '"':
                    if j + 1 < n and f[j + 1] == '"':
                        j += 2
                        continue
                    break
                j += 1
            for k in range(i, min(j + 1, n)):
                chars[k] = " "
            i = j + 1
            continue
        if c == "'":
            j = i + 1
            while j < n:
                if f[j] == "'":
                    if j + 1 < n and f[j + 1] == "'":
                        j += 2
                        continue
                    break
                j += 1
            sheet_spans.append((i, j))
            i = j + 1
            continue
        i += 1
    return "".join(chars), sheet_spans


def check_balance(masked: str) -> str | None:
    """括号配对与引号闭合。输入应为掩码串（字符串已空格化）。"""
    depth = 0
    for ch in masked:
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
            if depth < 0:
                return "右括号多于左括号"
    if depth != 0:
        return f"有 {depth} 个左括号未闭合"
    if masked.count("'") % 2 != 0:
        return "单引号（表名）未闭合"
    return None


def resolve_sheet_prefix(masked: str, pos: int) -> tuple[str | None, bool]:
    """查 pos 之前是否紧邻 'Sheet!' 前缀。返回 (表名, 是否外部引用)。"""
    m = PREFIX_RE.search(masked[:pos])
    if not m:
        return None, False
    name = m.group(2) if m.group(2) is not None else m.group(3)
    if name is not None:
        name = name.replace("''", "'")
    return name, m.group(1) is not None


def iter_refs(f: str, current_sheet: str, masked: str | None = None,
              sheet_spans: list[tuple[int, int]] | None = None):
    """产出公式中的引用：(sheet, c1, r1, c2, r2, external)。

    区间引用 A1:B5 的 sheet 语义与 Excel 一致——两端同表。
    """
    if masked is None:
        masked, sheet_spans = mask_formula(f)
    spans = sheet_spans or []
    matches = []
    for m in CELL_RE.finditer(masked):
        if any(a <= m.start() <= b for a, b in spans):
            continue
        matches.append(m)
    i = 0
    while i < len(matches):
        m = matches[i]
        try:
            c1, r1 = parse_a1(m.group(1))
        except ValueError:
            i += 1
            continue
        sheet, external = resolve_sheet_prefix(masked, m.start())
        sheet = sheet or current_sheet
        c2 = r2 = None
        if i + 1 < len(matches):
            gap = masked[m.end():matches[i + 1].start()]
            if gap == ":":
                try:
                    c2, r2 = parse_a1(matches[i + 1].group(1))
                except ValueError:
                    c2 = r2 = None
                else:
                    i += 1
        yield sheet, c1, r1, c2, r2, external
        i += 1


def extract_functions(masked: str) -> list[str]:
    return [m.group(1).upper() for m in FUNC_RE.finditer(masked)]


MAGIC_INT_EXEMPT = 12  # 绝对值 ≤ 此值的整数视为结构性常量（×2、+1、OFFSET 偏移等）


def extract_magic_numbers(masked: str) -> list[str]:
    """找出可疑的魔数字面量。

    豁免绝对值 ≤ MAGIC_INT_EXEMPT 的整数：这类几乎都是结构性常量。
    xlsx 技能所指的假设（增长率/利润率/倍数）通常是非整数，故小数一律不豁免。
    """
    tmp = FUNC_RE.sub(" ", masked)
    tmp = CELL_RE.sub(" ", tmp)
    out = []
    for m in NUM_RE.finditer(tmp):
        txt = m.group(1)
        try:
            val = float(txt)
        except ValueError:
            continue
        if val.is_integer() and abs(val) <= MAGIC_INT_EXEMPT:
            continue
        out.append(txt)
    return out


# ---------------------------------------------------------------- 结构校验


class Problems:
    def __init__(self):
        self.errors: list[dict] = []
        self.warnings: list[dict] = []

    def error(self, sheet, cell, rule, message, formula=None):
        item = {"sheet": sheet, "cell": cell, "rule": rule, "message": message}
        if formula:
            item["formula"] = formula
        self.errors.append(item)

    def warn(self, sheet, cell, rule, message, formula=None):
        item = {"sheet": sheet, "cell": cell, "rule": rule, "message": message}
        if formula:
            item["formula"] = formula
        self.warnings.append(item)


def validate_structure(obj, prob: Problems) -> bool:
    """结构校验（硬约束 §7）。返回是否有致命结构错误。"""
    ok = True
    if not isinstance(obj, dict):
        prob.error("-", "-", "ROOT_TYPE", "顶层必须是 JSON 对象")
        return False
    if obj.get("$schema") != SCHEMA_ID:
        prob.error("-", "-", "SCHEMA_ID",
                   f'$schema 必须是 "{SCHEMA_ID}"，实际为 {obj.get("$schema")!r}')
        ok = False
    conv = obj.get("conventions")
    if conv is not None and conv != "financial-model":
        prob.error("-", "-", "CONVENTIONS", f"conventions 只允许 \"financial-model\"，实际为 {conv!r}")
        ok = False
    sheets = obj.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        prob.error("-", "-", "SHEETS", "sheets 必须是非空数组")
        return False

    seen = set()
    for si, sh in enumerate(sheets):
        tag = sh.get("name") if isinstance(sh, dict) else f"#{si}"
        if not isinstance(sh, dict):
            prob.error(str(tag), "-", "SHEET_TYPE", "sheet 必须是对象")
            ok = False
            continue
        extra = set(sh) - SHEET_KEYS
        if extra:
            prob.error(tag, "-", "SHEET_KEYS", f"sheet 含未知字段 {sorted(extra)}")
            ok = False
        name = sh.get("name")
        if not isinstance(name, str) or not name:
            prob.error(str(tag), "-", "SHEET_NAME", "name 必须是非空字符串")
            ok = False
        else:
            if len(name) > 31:
                prob.error(name, "-", "SHEET_NAME", f"name 超过 31 字符（{len(name)}）")
                ok = False
            bad = BAD_SHEET_CHARS & set(name)
            if bad:
                prob.error(name, "-", "SHEET_NAME", f"name 含非法字符 {sorted(bad)}")
                ok = False
            if name in seen:
                prob.error(name, "-", "SHEET_NAME", "name 重复")
                ok = False
            seen.add(name)

        header = sh.get("header", True)
        if not isinstance(header, bool):
            prob.error(str(tag), "-", "HEADER", "header 必须是布尔值")
            ok = False
        columns = sh.get("columns")
        ncol = None
        if header:
            if columns is None:
                prob.error(str(tag), "-", "COLUMNS", "header 为 true 时必须提供 columns")
                ok = False
            elif not isinstance(columns, list) or not all(isinstance(c, str) for c in columns):
                prob.error(str(tag), "-", "COLUMNS", "columns 必须是字符串数组")
                ok = False
            else:
                ncol = len(columns)
        elif columns is not None:
            if not isinstance(columns, list):
                prob.error(str(tag), "-", "COLUMNS", "columns 必须是数组")
                ok = False
            else:
                ncol = len(columns)

        rows = sh.get("rows")
        if not isinstance(rows, list):
            prob.error(str(tag), "-", "ROWS", "rows 必须是数组")
            ok = False
            rows = []
        if ncol is not None:
            for ri, row in enumerate(rows):
                if not isinstance(row, list):
                    prob.error(name, f"第{ri + 1}数据行", "ROW_TYPE", "数据行必须是数组")
                    ok = False
                    continue
                if len(row) != ncol:
                    prob.error(name, f"第{ri + 1}数据行", "ROW_WIDTH",
                               f"期望 {ncol} 列，实际 {len(row)} 列")
                    ok = False

        if not isinstance(sh.get("merges", []), list):
            prob.error(name, "-", "MERGES", "merges 必须是数组")
            ok = False
        cw = sh.get("columnWidths", {})
        if not isinstance(cw, dict):
            prob.error(name, "-", "WIDTHS", "columnWidths 必须是对象")
            ok = False

        for ri, row in enumerate(rows):
            if not isinstance(row, list):
                continue
            for ci, cell in enumerate(row):
                if not validate_cell(cell, name, ri, ci, prob):
                    ok = False
    return ok


def validate_cell(cell, sheet: str, ri: int, ci: int, prob: Problems) -> bool:
    """校验单个单元格。返回是否合法。"""
    coord = f"第{ri + 1}数据行第{ci + 1}列"

    if cell is None:
        return True
    if isinstance(cell, str):
        if cell == "":
            prob.warn(sheet, coord, "EMPTY_STRING",
                      '空字符串 "" 经 openpyxl 写入后回读为 null（写出的是 '
                      '<c t="inlineStr" />，无内容）；要表示空单元格请直接用 null')
        return True
    if isinstance(cell, bool):
        return True
    if is_number(cell):
        if float(cell).is_integer() and abs(cell) >= 10 ** 15:
            prob.warn(sheet, coord, "PRECISION",
                      "整数超过 15 位有效数字，建议改用字符串以免丢精度")
        return True
    if not isinstance(cell, dict):
        prob.error(sheet, coord, "CELL_TYPE", f"不支持的单元格类型 {type(cell).__name__}")
        return False

    ok = True
    extra = set(cell) - CELL_OBJ_KEYS
    if extra:
        prob.error(sheet, coord, "CELL_KEYS", f"单元格含未知字段 {sorted(extra)}")
        ok = False
    if not ({"v", "f", "link"} & set(cell)):
        prob.error(sheet, coord, "CELL_EMPTY", "对象单元格必须至少含 v / f / link 之一")
        ok = False
    if "preset" in cell and "fmt" in cell:
        prob.error(sheet, coord, "PRESET_FMT", "preset 与 fmt 互斥")
        ok = False
    if "preset" in cell and cell["preset"] not in PRESETS:
        prob.error(sheet, coord, "PRESET_UNKNOWN",
                   f"未知 preset {cell['preset']!r}，可选：{sorted(PRESETS)}")
        ok = False
    if "t" in cell and cell["t"] not in T_DEFAULT_FORMAT:
        prob.error(sheet, coord, "T_UNKNOWN", f"t 只能是 {sorted(T_DEFAULT_FORMAT)}")
        ok = False
    if "t" in cell and "v" not in cell:
        prob.error(sheet, coord, "T_NO_VALUE", "t 必须与 v 同时出现")
        ok = False
    roles = cell.get("role")
    if roles is not None:
        role_list = [roles] if isinstance(roles, str) else roles
        if not isinstance(role_list, list) or not all(isinstance(r, str) for r in role_list):
            prob.error(sheet, coord, "ROLE_TYPE", "role 必须是字符串或字符串数组")
            ok = False
        else:
            unk = [r for r in role_list if r not in ROLE_NAMES]
            if unk:
                prob.error(sheet, coord, "ROLE_UNKNOWN",
                           f"未知 role {unk}，可选：{sorted(ROLE_NAMES)}")
                ok = False
    if "f" in cell:
        if not isinstance(cell["f"], str) or not cell["f"].strip():
            prob.error(sheet, coord, "FORMULA_EMPTY", "f 必须是非空字符串")
            ok = False
    return ok


# ---------------------------------------------------------------- 公式静态检查


def check_formulas(obj, prob: Problems, *, extra_funcs: set[str]) -> dict:
    """公式静态检查（§4.2）。返回统计信息。"""
    sheets = obj.get("sheets", [])
    sheet_by_name = {s.get("name"): s for s in sheets if isinstance(s, dict)}
    used = {}
    formula_cells: dict[tuple[str, int, int], str] = {}

    def excel_coord(sh, ri, ci):
        r = ri + (2 if sh.get("header", True) else 1)
        return f"{get_column_letter(ci + 1)}{r}"

    # 收集网格尺寸与公式单元格
    for sh in sheets:
        if not isinstance(sh, dict):
            continue
        name = sh.get("name")
        if not isinstance(name, str) or not name:
            continue
        rows = sh.get("rows") or []
        header = sh.get("header", True)
        cols = sh.get("columns") or []
        ncol = max([len(cols)] + [len(r) for r in rows if isinstance(r, list)] + [0])
        nrow = len(rows) + (1 if header else 0)
        used[name] = (nrow, ncol)
        for ri, row in enumerate(rows):
            if not isinstance(row, list):
                continue
            for ci, cell in enumerate(row):
                if isinstance(cell, dict) and isinstance(cell.get("f"), str):
                    formula_cells[(name, ci + 1, ri + (2 if header else 1))] = cell["f"]

    stats = {"sheets": len(sheets), "formulas": len(formula_cells),
             "cached_errors": 0, "unverified": len(formula_cells)}

    whitelist = EXCEL_FUNCTIONS | {f.upper() for f in extra_funcs}

    for sh in sheets:
        if not isinstance(sh, dict):
            continue
        name = sh.get("name")
        if not isinstance(name, str) or not name:
            continue
        header = sh.get("header", True)
        for ri, row in enumerate(sh.get("rows") or []):
            if not isinstance(row, list):
                continue
            for ci, cell in enumerate(row):
                if not isinstance(cell, dict) or not isinstance(cell.get("f"), str):
                    continue
                f = cell["f"]
                coord = excel_coord(sh, ri, ci)
                body = f[1:] if f.startswith("=") else f

                if not body.strip():
                    prob.error(name, coord, "FORMULA_EMPTY", "公式为空", f)
                    continue

                masked, spans = mask_formula(body)

                if (bal := check_balance(masked)) is not None:
                    prob.error(name, coord, "SYNTAX", f"括号/引号不配对：{bal}", f)

                for lit in FATAL_ERROR_LITERALS:
                    if lit in masked:
                        prob.error(name, coord, "ERROR_LITERAL",
                                   f"公式含错误字面量 {lit}", f)
                        break

                for fn in extract_functions(masked):
                    if fn not in whitelist:
                        prob.error(name, coord, "NAME_UNKNOWN",
                                   f"未知函数 {fn}，可能拼写错误", f)

                external = False
                for ref_sheet, c1, r1, c2, r2, ext in iter_refs(body, name, masked, spans):
                    if ext:
                        external = True
                        continue
                    if ref_sheet not in sheet_by_name:
                        prob.error(name, coord, "REF_SHEET",
                                   f"引用了不存在的工作表「{ref_sheet}」", f)
                        continue
                    for cc, rr, tag in ((c1, r1, "起始"), (c2, r2, "结束")):
                        if cc is None:
                            continue
                        if cc > MAX_COL or rr > MAX_ROW:
                            prob.error(name, coord, "REF_BOUNDS",
                                       f"{tag}引用越界（{get_column_letter(cc) if cc <= MAX_COL else 'XFE+'}{rr}）", f)

                magic = extract_magic_numbers(masked)
                if magic:
                    prob.warn(name, coord, "MAGIC_NUMBER",
                              f"公式含魔数 {'、'.join(sorted(set(magic)))}，"
                              f"建议改用单元格引用", f)

    detect_cycles(formula_cells, used, prob)
    return stats


def detect_cycles(formula_cells, used, prob: Problems):
    """依赖图找环。区间引用按已用范围裁剪后展开。"""
    edges: dict[tuple[str, int, int], set] = {}
    for (sheet, col, row), f in formula_cells.items():
        body = f[1:] if f.startswith("=") else f
        masked, spans = mask_formula(body)
        deps = set()
        for ref_sheet, c1, r1, c2, r2, ext in iter_refs(body, sheet, masked, spans):
            if ext or ref_sheet not in used:
                continue
            urow, ucol = used[ref_sheet]
            c_lo, c_hi = sorted((c1, c2 if c2 is not None else c1))
            r_lo, r_hi = sorted((r1, r2 if r2 is not None else r1))
            c_lo, c_hi = max(c_lo, 1), min(c_hi, max(ucol, 1))
            r_lo, r_hi = max(r_lo, 1), min(r_hi, max(urow, 1))
            if (c_hi - c_lo + 1) * (r_hi - r_lo + 1) > 200_000:
                continue
            for cc in range(c_lo, c_hi + 1):
                for rr in range(r_lo, r_hi + 1):
                    if (ref_sheet, cc, rr) in formula_cells:
                        deps.add((ref_sheet, cc, rr))
            if len(deps) > 200_000:
                break
        edges[(sheet, col, row)] = deps

    WHITE, GREY, BLACK = 0, 1, 2
    color = {k: WHITE for k in edges}
    stack: list = []
    seen_cycles: set[frozenset] = set()

    def dfs(node):
        color[node] = GREY
        stack.append(node)
        for nxt in edges.get(node, ()):
            if nxt not in color:
                continue
            if color[nxt] == GREY:
                # 不变式：所有 GREY 节点都在 stack 上
                i = stack.index(nxt)
                cycle = stack[i:] + [nxt]
                key = frozenset(cycle)
                if key not in seen_cycles:
                    seen_cycles.add(key)
                    cells = " → ".join(
                        f"{s}!{get_column_letter(c)}{r}" for s, c, r in cycle)
                    prob.error(node[0], f"{get_column_letter(node[1])}{node[2]}",
                               "CIRCULAR", f"循环引用：{cells}")
            elif color[nxt] == WHITE:
                dfs(nxt)
        stack.pop()
        color[node] = BLACK

    for node in list(edges):
        if color[node] == WHITE:
            dfs(node)
    stack.clear()


# ---------------------------------------------------------------- 读取 xlsx


def read_bounds(ws, keep_empty: bool):
    if keep_empty:
        return max(ws.max_row or 1, 1), max(ws.max_column or 1, 1)
    max_r = max_c = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                max_r = max(max_r, cell.row)
                max_c = max(max_c, cell.column)
    return max_r, max_c


def role_of(cell) -> list[str]:
    roles = []
    try:
        color = cell.font.color
        if color is not None and getattr(color, "type", None) == "rgb" and color.rgb:
            rgb = str(color.rgb)[-6:].upper()
            if rgb in FONT_ROLE_BY_RGB:
                roles.append(FONT_ROLE_BY_RGB[rgb])
    except Exception:
        pass
    try:
        fill = cell.fill
        if fill is not None and fill.fill_type == "solid":
            sc = fill.start_color
            if sc is not None and getattr(sc, "type", None) == "rgb" and sc.rgb:
                rgb = str(sc.rgb)[-6:].upper()
                if rgb in FILL_ROLE_BY_RGB:
                    roles.append(FILL_ROLE_BY_RGB[rgb])
    except Exception:
        pass
    return roles


def format_has_time(number_format: str | None) -> bool:
    """Excel 不区分日期/日期时间，靠数字格式判断。m 有月份/分钟歧义，故只看 h/s。"""
    if not number_format:
        return False
    nf = re.sub(r'"[^"]*"', "", number_format)   # 去字面量
    nf = re.sub(r"\[[^\]]*\]", "", nf)           # 去 [h] / 颜色 等
    nf = re.sub(r"\\.", "", nf)                  # 去转义
    return bool(re.search(r"[hs]", nf, re.I))


def read_cell(fws, vws, row, col, opts) -> object:
    fc = fws.cell(row=row, column=col)
    raw = fc.value
    cached = vws.cell(row=row, column=col).value if vws is not None else None

    if isinstance(raw, ArrayFormula):
        out: dict = {"f": str(raw.text).lstrip("="), "array": True}
        if cached is not None and not opts["no_cached"]:
            out["v"] = cached
        return _decorate(out, fc, opts)

    if isinstance(raw, str) and raw.startswith("="):
        out = {"f": raw[1:]}
        if cached is not None and not opts["no_cached"]:
            out["v"] = cached
        return _decorate(out, fc, opts)

    if isinstance(raw, (dt.datetime, dt.date, dt.time)):
        if isinstance(raw, dt.datetime):
            if format_has_time(fc.number_format):
                t, s = "datetime", raw.isoformat(sep=" ")
            else:
                t, s = "date", raw.date().isoformat()
        elif isinstance(raw, dt.date):
            t, s = "date", raw.isoformat()
        else:
            t, s = "time", raw.isoformat()
        return _decorate({"v": s, "t": t}, fc, opts)

    return _decorate(raw, fc, opts)


def _decorate(value, cell, opts):
    """附加 preset / fmt / role / note。”"""
    extra: dict = {}
    if not opts["no_fmt"]:
        nf = cell.number_format
        if nf and nf != "General":
            if nf in PRESET_BY_FORMAT:
                extra["preset"] = PRESET_BY_FORMAT[nf]
            else:
                extra["fmt"] = nf
    if not opts["no_role"]:
        roles = role_of(cell)
        if roles:
            extra["role"] = roles[0] if len(roles) == 1 else roles
    try:
        if cell.comment is not None and cell.comment.text:
            extra["note"] = cell.comment.text
    except Exception:
        pass
    if not extra:
        return value
    if isinstance(value, dict):
        value.update(extra)
        return value
    return {"v": value, **extra}


def xlsx_to_obj(path: Path, opts) -> dict:
    formulas = load_workbook(path, data_only=False)
    try:
        values = load_workbook(path, data_only=True)
    except Exception:
        values = None
    if opts.get("no_cached"):
        values = None

    wanted = opts.get("sheets")
    sheets_out = []
    for ws in formulas.worksheets:
        if wanted and ws.title not in wanted:
            continue
        vws = values[ws.title] if values is not None and ws.title in values.sheetnames else None
        header = not opts["no_header"]
        max_r, max_c = read_bounds(ws, opts["keep_empty"])

        columns = []
        if header and max_r >= 1 and max_c >= 1:
            for c in range(1, max_c + 1):
                v = ws.cell(row=1, column=c).value
                columns.append("" if v is None else str(v))

        sh: dict = {"name": ws.title, "header": header}
        if header:
            sh["columns"] = columns

        rows = []
        start = 2 if header else 1
        for r in range(start, max_r + 1):
            row = [read_cell(ws, vws, r, c, opts) for c in range(1, max_c + 1)]
            rows.append(row)
        sh["rows"] = rows

        if ws.merged_cells.ranges:
            sh["merges"] = sorted(str(x) for x in ws.merged_cells.ranges)
        widths = {}
        for letter, dim in ws.column_dimensions.items():
            if dim.width:
                w = round(float(dim.width), 2)
                widths[letter] = int(w) if float(w).is_integer() else w
        if widths:
            sh["columnWidths"] = widths
        if ws.freeze_panes:
            sh["freeze"] = str(ws.freeze_panes)
        sheets_out.append(sh)

    formulas.close()
    if values is not None:
        values.close()

    if wanted:
        missing = [s for s in wanted if s not in {x["name"] for x in sheets_out}]
        if missing:
            die(f"工作表不存在：{missing}")
    if not sheets_out:
        die("没有任何工作表可导出")

    obj = {"$schema": SCHEMA_ID, "sheets": sheets_out}
    return obj


# ---------------------------------------------------------------- 写 xlsx


def parse_temporal(value, t: str):
    if not isinstance(value, str):
        return value
    try:
        if t == "date":
            return dt.date.fromisoformat(value)
        if t == "datetime":
            return dt.datetime.fromisoformat(value)
        return dt.time.fromisoformat(value)
    except ValueError:
        try:
            return dt.datetime.fromisoformat(value)
        except ValueError:
            die(f"无法解析 {t} 值：{value!r}")


def apply_style(cell, roles, *, bold=False):
    for role in roles:
        font_rgb, fill_rgb = ROLE_STYLE[role]
        if font_rgb:
            f = cell.font
            cell.font = Font(name=f.name, size=f.size, bold=f.bold or bold,
                             italic=f.italic, color=font_rgb)
        if fill_rgb:
            cell.fill = PatternFill("solid", start_color=fill_rgb)
    if bold and not roles:
        f = cell.font
        cell.font = Font(name=f.name, size=f.size, bold=True)


def write_cell(ws, coord, spec, ctx):
    cell = ws[coord]
    roles: list[str] = []
    wrap = False

    if isinstance(spec, dict):
        raw_roles = spec.get("role")
        if raw_roles:
            roles = [raw_roles] if isinstance(raw_roles, str) else list(raw_roles)
        if ctx["conventions"] == "financial-model" and not roles:
            roles = infer_roles(spec, ctx["sheet_names"])

        if "f" in spec:
            body = spec["f"]
            body = body[1:] if body.startswith("=") else body
            if spec.get("array"):
                cell.value = ArrayFormula(ref=coord, text="=" + body)
            else:
                cell.value = "=" + body
            # 缓存值不落盘：openpyxl 只会写出空的 <v/>（见规范 §4.3）
        elif "t" in spec:
            cell.value = parse_temporal(spec.get("v"), spec["t"])
        elif "v" in spec:
            cell.value = spec["v"]

        if "link" in spec:
            cell.hyperlink = spec["link"]
            if "v" not in spec:
                cell.value = spec["link"]

        if "preset" in spec:
            cell.number_format = PRESETS[spec["preset"]]
        elif "fmt" in spec:
            cell.number_format = spec["fmt"]
        elif "t" in spec:
            cell.number_format = T_DEFAULT_FORMAT[spec["t"]]

        if "note" in spec:
            cell.comment = Comment(spec["note"], "sheet-json")
    else:
        cell.value = spec

    if isinstance(cell.value, str) and "\n" in cell.value and not ctx["no_wrap"]:
        wrap = True
    if wrap:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    return roles


def infer_roles(spec: dict, sheet_names: set[str]) -> list[str]:
    f = spec.get("f")
    if f:
        masked, _ = mask_formula(f)
        if re.search(r"\[[^\]]*\]", masked):
            return ["external"]
        for m in re.finditer(r"([^!'()+\-*/^&<>=,;:\[\]\s]+)!", masked):
            if m.group(1).strip("'") in sheet_names:
                return ["link"]
        return ["formula"]
    if is_number(spec.get("v")):
        return ["input"]
    return []


def obj_to_workbook(obj, opts) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    ctx = {
        "conventions": obj.get("conventions"),
        "sheet_names": {s.get("name") for s in obj.get("sheets", [])},
        "no_wrap": opts["no_wrap"],
    }
    for sh in obj.get("sheets", []):
        ws = wb.create_sheet(title=sh["name"])
        header = sh.get("header", True)
        columns = sh.get("columns") or []
        rows = sh.get("rows") or []

        if header and columns:
            for ci, name in enumerate(columns, start=1):
                c = ws.cell(row=1, column=ci, value=name)
                apply_style(c, [], bold=True)

        start = 2 if header else 1
        for ri, row in enumerate(rows):
            r = start + ri
            for ci, spec in enumerate(row):
                if spec is None:
                    continue
                coord = f"{get_column_letter(ci + 1)}{r}"
                roles = write_cell(ws, coord, spec, ctx)
                if roles:
                    apply_style(ws[coord], roles)

        for rng in sh.get("merges") or []:
            ws.merge_cells(rng)
        for letter, width in (sh.get("columnWidths") or {}).items():
            ws.column_dimensions[letter.upper()].width = float(width)
        if sh.get("freeze"):
            ws.freeze_panes = sh["freeze"]
    if not wb.sheetnames:
        wb.create_sheet(title="Sheet1")
    return wb


# ---------------------------------------------------------------- 缓存错误扫描


def scan_cached_errors(path: Path) -> list[dict]:
    """读已保存文件里的错误缓存值。无需计算引擎——错误值是算好的结果。"""
    out = []
    try:
        wb = load_workbook(path, data_only=True)
    except Exception:
        return out
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v in EXCEL_ERRORS:
                    out.append({"sheet": ws.title, "cell": cell.coordinate,
                                "rule": "CACHED_ERROR", "message": f"缓存值为 {v}"})
    wb.close()
    return out


# ---------------------------------------------------------------- 各子命令


def build_report(obj, *, extra_funcs: set[str], strict: bool,
                 cached_errors: list[dict] | None = None) -> dict:
    prob = Problems()
    validate_structure(obj, prob)
    # 公式检查独立于结构校验：结构有错时也要把公式问题一并报出，
    # 否则用户只能看到第一批错误，修完再跑才发现还有下一批。
    stats = check_formulas(obj, prob, extra_funcs=extra_funcs)
    if cached_errors:
        prob.errors.extend(cached_errors)
        stats["cached_errors"] = len(cached_errors)
    if strict:
        prob.errors.extend(prob.warnings)
        prob.warnings = []
    status = "errors_found" if prob.errors else ("warnings" if prob.warnings else "success")
    return {"status": status, "errors": prob.errors, "warnings": prob.warnings, "stats": stats}


def load_json(path: Path) -> dict:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        die(f"文件不存在：{path}")
    except json.JSONDecodeError as e:
        die(f"JSON 解析失败（第 {e.lineno} 行第 {e.colno} 列）：{e.msg}")


def cmd_to_json(args) -> int:
    src = Path(args.input)
    if not src.exists():
        die(f"文件不存在：{src}")
    opts = {"sheets": args.sheet, "no_cached": args.no_cached, "no_header": args.no_header,
            "no_fmt": args.no_fmt, "no_role": args.no_role, "keep_empty": args.keep_empty}
    obj = xlsx_to_obj(src, opts)
    out = Path(args.output) if args.output else src.with_suffix(".json")
    out.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    ncell = sum(len(r) for s in obj["sheets"] for r in s["rows"])
    print(f"✓ {src} → {out}（{len(obj['sheets'])} 个工作表，{ncell} 个单元格）")
    return 0


def cmd_check(args) -> int:
    src = Path(args.input)
    if not src.exists():
        die(f"文件不存在：{src}")
    cached = None
    if src.suffix.lower() in (".xlsx", ".xlsm"):
        obj = xlsx_to_obj(src, {"sheets": None, "no_cached": False, "no_header": False,
                                "no_fmt": True, "no_role": True, "keep_empty": False})
        cached = scan_cached_errors(src)
    else:
        obj = load_json(src)
    report = build_report(obj, extra_funcs=set(args.extra_funcs or []),
                          strict=args.strict, cached_errors=cached)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 1 if report["status"] == "errors_found" else 0


def cmd_to_excel(args) -> int:
    src = Path(args.input)
    obj = load_json(src)
    report = build_report(obj, extra_funcs=set(args.extra_funcs or []), strict=args.strict)
    if report["status"] == "errors_found" and not args.force:
        print("校验未通过，未生成文件。加 --force 可强制写出。", file=sys.stderr)
        print(json.dumps(report, ensure_ascii=False, indent=2), file=sys.stderr)
        return 1

    if args.pad:
        for sh in obj.get("sheets", []):
            cols = sh.get("columns") or []
            if sh.get("header", True) and cols:
                for row in sh.get("rows", []):
                    while len(row) < len(cols):
                        row.append(None)

    wb = obj_to_workbook(obj, {"no_wrap": args.no_wrap})
    out = Path(args.output) if args.output else src.with_suffix(".xlsx")
    wb.save(out)
    ncell = sum(len(r) for s in obj["sheets"] for r in s["rows"])
    print(f"✓ {src} → {out}（{len(obj['sheets'])} 个工作表，{ncell} 个单元格）")
    if report["warnings"]:
        print(f"  提示：{len(report['warnings'])} 条 warning（--strict 可升级为错误）", file=sys.stderr)
    u = report["stats"].get("unverified", 0)
    if u:
        print(f"  注意：{u} 个公式的真值未被验证（无计算引擎，规范 §4.2）", file=sys.stderr)
    return 0


def flatten(obj) -> dict:
    out = {}
    for sh in obj.get("sheets", []):
        header = sh.get("header", True)
        cols = sh.get("columns") or []
        cells = {}
        if header:
            for ci, name in enumerate(cols):
                cells[(1, ci + 1)] = ("value", name)
        start = 2 if header else 1
        for ri, row in enumerate(sh.get("rows") or []):
            for ci, cell in enumerate(row):
                key = (start + ri, ci + 1)
                if isinstance(cell, dict):
                    if "f" in cell:
                        body = cell["f"]
                        body = body[1:] if body.startswith("=") else body
                        cells[key] = ("formula", re.sub(r"\s+", "", body))
                    else:
                        cells[key] = ("value", cell.get("v"))
                else:
                    cells[key] = ("value", cell)
        out[sh.get("name")] = cells
    return out


def same_value(a, b, tol) -> bool:
    if is_number(a) and is_number(b):
        return abs(float(a) - float(b)) <= tol
    if isinstance(a, bool) != isinstance(b, bool):
        return False
    return a == b


def compare(orig: dict, rt: dict, tol: float) -> list[dict]:
    fa, fb = flatten(orig), flatten(rt)
    diffs = []
    for name in sorted(set(fa) | set(fb)):
        if name not in fb:
            diffs.append({"sheet": name, "cell": "-", "kind": "sheet_missing", "original": name})
            continue
        if name not in fa:
            diffs.append({"sheet": name, "cell": "-", "kind": "sheet_extra", "roundtrip": name})
            continue
        ca, cb = fa[name], fb[name]
        for key in sorted(set(ca) | set(cb)):
            coord = f"{get_column_letter(key[1])}{key[0]}"
            if key not in cb:
                diffs.append({"sheet": name, "cell": coord, "kind": "missing",
                              "original": ca[key][1]})
            elif key not in ca:
                diffs.append({"sheet": name, "cell": coord, "kind": "extra",
                              "roundtrip": cb[key][1]})
            else:
                (ka, va), (kb, vb) = ca[key], cb[key]
                if ka != kb:
                    diffs.append({"sheet": name, "cell": coord, "kind": "kind_changed",
                                  "original": f"{ka}:{va}", "roundtrip": f"{kb}:{vb}"})
                elif ka == "formula":
                    if va != vb:
                        diffs.append({"sheet": name, "cell": coord, "kind": "formula",
                                      "original": va, "roundtrip": vb})
                elif not same_value(va, vb, tol):
                    diffs.append({"sheet": name, "cell": coord, "kind": "value",
                                  "original": va, "roundtrip": vb})
    return diffs


def cmd_verify(args) -> int:
    src = Path(args.input)
    if not src.exists():
        die(f"文件不存在：{src}")
    opts = {"sheets": args.sheet, "no_cached": False, "no_header": False,
            "no_fmt": True, "no_role": True, "keep_empty": args.keep_empty}
    orig = xlsx_to_obj(src, opts)
    wb = obj_to_workbook(orig, {"no_wrap": False})
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td) / "roundtrip.xlsx"
        wb.save(tmp)
        rt = xlsx_to_obj(tmp, opts)
    diffs = compare(orig, rt, args.tolerance)
    ncell = sum(len(c) for c in flatten(orig).values())
    report = {
        "status": "different" if diffs else "identical",
        "sheets_compared": len(orig.get("sheets", [])),
        "cells_compared": ncell,
        "difference_count": len(diffs),
        "differences": diffs[:50],
        "note": "公式缓存值不参与比较（openpyxl 写路径不落盘，规范 §4.3）",
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 1 if diffs else 0


# ---------------------------------------------------------------- CLI


def main(argv=None) -> int:
    p = argparse.ArgumentParser(
        prog="sheetjson",
        description="sheet-json ⇄ Excel 双向转换与静态校验（仅依赖 openpyxl）")
    sub = p.add_subparsers(dest="cmd", required=True)

    a = sub.add_parser("to-json", help="xlsx → sheet-json")
    a.add_argument("input")
    a.add_argument("-o", "--output")
    a.add_argument("--sheet", action="append", help="只导出指定工作表，可重复")
    a.add_argument("--no-cached", action="store_true", help="不输出公式缓存值")
    a.add_argument("--no-header", action="store_true", help="首行也作为数据")
    a.add_argument("--no-fmt", action="store_true", help="不输出 preset/fmt")
    a.add_argument("--no-role", action="store_true", help="不输出 role")
    a.add_argument("--keep-empty", action="store_true", help="保留尾部空行空列")
    a.set_defaults(func=cmd_to_json)

    b = sub.add_parser("to-excel", help="sheet-json → xlsx（写入前静态校验）")
    b.add_argument("input")
    b.add_argument("-o", "--output")
    b.add_argument("--pad", action="store_true", help="列数不足的行用 null 补齐")
    b.add_argument("--no-wrap", action="store_true", help="不自动开启单元格换行")
    b.add_argument("--strict", action="store_true", help="warning 升级为 error")
    b.add_argument("--force", action="store_true", help="校验失败仍强制写出")
    b.add_argument("--extra-funcs", action="append", help="追加函数名白名单")
    b.set_defaults(func=cmd_to_excel)

    c = sub.add_parser("check", help="静态校验，输出 JSON 报告")
    c.add_argument("input", help=".json 或 .xlsx")
    c.add_argument("--strict", action="store_true", help="warning 升级为 error")
    c.add_argument("--extra-funcs", action="append", help="追加函数名白名单")
    c.set_defaults(func=cmd_check)

    d = sub.add_parser("verify", help="xlsx → json → xlsx 往返比对")
    d.add_argument("input")
    d.add_argument("--sheet", action="append")
    d.add_argument("--tolerance", type=float, default=1e-9)
    d.add_argument("--keep-empty", action="store_true")
    d.set_defaults(func=cmd_verify)

    args = p.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
